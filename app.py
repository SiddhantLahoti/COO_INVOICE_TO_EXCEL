import io
import re
from collections import OrderedDict
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from pdfminer.high_level import extract_pages
from pdfminer.layout import LTTextContainer, LTTextLineHorizontal


st.set_page_config(page_title="Invoice to Excel Converter", page_icon="📄", layout="wide")


def parse_pdf_stream(pdf_file):
    """Parses a PDF file buffer directly from Streamlit uploader."""
    lines = []
    # Read bytes into memory buffer
    pdf_bytes = io.BytesIO(pdf_file.read())
    
    pages = list(extract_pages(pdf_bytes))

    if not pages:
        return None

    page_1 = pages[0]
    last_page = pages[-1]
    
    for element in page_1:
        if isinstance(element, LTTextContainer):
            for text_line in element:
                if isinstance(text_line, LTTextLineHorizontal):
                    t = text_line.get_text().strip()
                    if t:
                        lines.append((text_line.bbox, t))

    if not lines:
        return None

    # 1. Broadened Invoice Number and Date Extraction
    invoice_no = "UNKNOWN"
    invoice_date = ""

    # Sort lines top-to-bottom
    lines_sorted = sorted(lines, key=lambda x: -x[0][1])
    conv_rate = 1.0

    # Extract full text of the last page to handle split text boxes
    last_page_text = ""
    for element in last_page:
        if isinstance(element, LTTextContainer):
            last_page_text += " " + element.get_text()

    # Search across the combined text
    m_conv = re.search(r'Conv\s*Rs\.?\s*[:\-]?\s*([0-9]+\.[0-9]+)', last_page_text, re.IGNORECASE)
    if m_conv:
        conv_rate = float(m_conv.group(1))
    for bbox, text in lines_sorted:
        # Matches PJ-26000784, MJ-26001420, etc.
        if invoice_no == "UNKNOWN":
            m_inv = re.search(r'\b[A-Za-z]{2,4}-\d+\b', text)
            if m_inv:
                invoice_no = m_inv.group(0).upper()

        # Matches dates like 22/09/2026 in the top invoice header
        if not invoice_date and bbox[1] > 600:
            m_date = re.search(r'\b\d{2}[/-]\d{2}[/-]\d{4}\b', text)
            if m_date:
                invoice_date = m_date.group(0)

    # Group lines by vertical Y position
    lines_sorted = sorted(lines, key=lambda x: -x[0][1])
    row_groups = []
    for bbox, text in lines_sorted:
        y0 = bbox[1]
        placed = False
        for g in row_groups:
            if abs(g['y'] - y0) < 3.5:
                g['items'].append((bbox[0], text))
                placed = True
                break
        if not placed:
            row_groups.append({'y': y0, 'items': [(bbox[0], text)]})

    for g in row_groups:
        g['items'].sort(key=lambda x: x[0])

    # HSN Code
    hsn_code = ""
    hsn_idx = -1
    for i, g in enumerate(row_groups):
        for x0, txt in g['items']:
            if re.match(r'^\d{8}$', txt):
                hsn_code = txt
                hsn_idx = i
                break
        if hsn_code:
            break

    # First Category Header
    first_cat_idx = -1
    if hsn_idx != -1:
        for i in range(hsn_idx + 1, len(row_groups)):
            row_str = " ".join([txt for _, txt in row_groups[i]['items']])
            if re.search(r'\b(PCS|PRS)\s*,', row_str):
                first_cat_idx = i
                break

    # Multi-line Description
    desc_lines = []
    if hsn_idx != -1 and first_cat_idx != -1:
        for i in range(hsn_idx + 1, first_cat_idx):
            line_str = " ".join([txt for _, txt in row_groups[i]['items']]).strip()
            if line_str and not line_str.startswith("(Gms)") and not line_str.startswith("NetWt"):
                desc_lines.append(line_str)

    base_description = " ".join(desc_lines).strip()

    # Category and Item rows
    groups = []
    current_category = None

    if first_cat_idx != -1:
        for i in range(first_cat_idx, len(row_groups)):
            g = row_groups[i]
            row_str = " ".join([txt for _, txt in g['items']])

            if any(marker in row_str for marker in ["RM KT", "Loss %", "NOTE :", "FOB US$"]):
                break

            if re.search(r'\b(PCS|PRS)\s*,', row_str):
                unit = "PCS" if "PCS" in row_str else "PRS"
                current_category = {
                    'header': row_str,
                    'unit': unit,
                    'has_plus': '+' in row_str,
                    'items': []
                }
                groups.append(current_category)
                continue

            if current_category is not None:
                item_name = None
                numbers = []
                for x0, txt in g['items']:
                    clean_num = txt.replace(',', '')
                    if re.match(r'^-?\d+(\.\d+)?$', clean_num):
                        numbers.append(float(clean_num))
                    else:
                        item_name = txt if item_name is None else f"{item_name} {txt}"

                if item_name and len(numbers) >= 3:
                    if len(numbers) >= 5:
                        net_wt, gross_wt, qty, _, amt = numbers[:5]
                    elif len(numbers) == 4:
                        net_wt, gross_wt, qty, amt = numbers
                    else:
                        net_wt = numbers[0]
                        gross_wt, qty, amt = numbers[0], numbers[1], numbers[2]

                    current_category['items'].append({
                        'name': item_name,
                        'net_wt': net_wt,      # <--- Add this
                        'gross_wt': gross_wt,
                        'qty': qty,
                        'amount': amt
                    })

    return {
    'invoice_no': invoice_no,
    'invoice_date': invoice_date,
    'hsn_code': hsn_code,
    'base_description': base_description,
    'categories': groups,
    'conv_rate': conv_rate  # <--- Added
}


def generate_excel(consolidated_data, invoices_list):
    """Creates formatted Excel workbook without KT column, matching manual structure."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidated Summary"

    headers = ["SR", "HSN", "DESCRIPTION", "GrossWt", "NET WT", "QTY", "AMT", "AMT INR"]
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 25

    header_font = Font(name="Calibri", size=11, bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 3
    start_data_row = 3
    data_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)

    sr = 1
    for key, data in consolidated_data.items():
        # Alternating blank row
        ws.cell(row=current_row, column=1, value=None)
        current_row += 1

        items_str = ", ".join(data['item_names'])
        kt_str = f" IN {', '.join(sorted(data['kts']))}" if data.get('kts') else ""
        base_description = f"{data.get('base_prefix', '')}{kt_str}"
        desc_text = f"{base_description}-{items_str}-{data['qty']} {data['unit']}"
        # SR
        c_sr = ws.cell(row=current_row, column=1, value=sr)
        c_sr.alignment = Alignment(horizontal="center")

        # HSN
        c_hsn = ws.cell(row=current_row, column=2, value=data['hsn'])
        c_hsn.alignment = Alignment(horizontal="center")

        # Description
        ws.cell(row=current_row, column=3, value=desc_text)

        # Col 4: GrossWt
        c_wt = ws.cell(row=current_row, column=4, value=round(data['gross_wt'], 3))
        c_wt.font = bold_font
        c_wt.number_format = '0.000'
        c_wt.alignment = Alignment(horizontal="right")

        # Col 5: NET WT
        c_net = ws.cell(row=current_row, column=5, value=round(data['net_wt'], 3))
        c_net.font = bold_font
        c_net.number_format = '0.000'
        c_net.alignment = Alignment(horizontal="right")

        # Col 6: QTY
        c_qty = ws.cell(row=current_row, column=6, value=data['qty'])
        c_qty.font = data_font
        c_qty.number_format = '#,##0'
        c_qty.alignment = Alignment(horizontal="right")

        # Col 7: AMT (USD)
        c_amt = ws.cell(row=current_row, column=7, value=round(data['amount'], 2))
        c_amt.font = bold_font
        c_amt.number_format = '#,##0.00'
        c_amt.alignment = Alignment(horizontal="right")

        # Col 8: AMT INR
        c_inr = ws.cell(row=current_row, column=8, value=round(data['amt_inr'], 2))
        c_inr.font = bold_font
        c_inr.number_format = '#,##0.00'
        c_inr.alignment = Alignment(horizontal="right")

        sr += 1
        current_row += 1

    # Blank row before total
    ws.cell(row=current_row, column=1, value=None)
    current_row += 1

    # Total Row with dynamic formulas
    end_data_row = current_row - 1
    total_lbl = ws.cell(row=current_row, column=3, value="TOTAL")
    total_lbl.font = bold_font
    total_lbl.alignment = Alignment(horizontal="right")

    # --- PASTE THE NEW CODE HERE ---
    # GrossWt (Column D / 4)
    c = ws.cell(row=current_row, column=4, value=f"=SUM(D{start_data_row}:D{end_data_row})")
    c.font = bold_font; c.alignment = Alignment(horizontal="right"); c.number_format = '0.000'

    # NetWt (Column E / 5)
    c = ws.cell(row=current_row, column=5, value=f"=SUM(E{start_data_row}:E{end_data_row})")
    c.font = bold_font; c.alignment = Alignment(horizontal="right"); c.number_format = '0.000'

    # QTY (Column F / 6)
    c = ws.cell(row=current_row, column=6, value=f"=SUM(F{start_data_row}:F{end_data_row})")
    c.font = bold_font; c.alignment = Alignment(horizontal="right"); c.number_format = '#,##0'

    # AMT (Column G / 7)
    c = ws.cell(row=current_row, column=7, value=f"=SUM(G{start_data_row}:G{end_data_row})")
    c.font = bold_font; c.alignment = Alignment(horizontal="right"); c.number_format = '#,##0.00'

    # AMT INR (Column H / 8)
    c = ws.cell(row=current_row, column=8, value=f"=SUM(H{start_data_row}:H{end_data_row})")
    c.font = bold_font; c.alignment = Alignment(horizontal="right"); c.number_format = '#,##0.00'
    # -------------------------------



    # --- ADD THIS BLOCK HERE ---
    # 2 blank rows below the TOTAL row
    current_row += 3  

    # Section Headers
    ws.cell(row=current_row, column=2, value="Invoice No.").font = bold_font
    ws.cell(row=current_row, column=3, value="Invoice Date").font = bold_font
    current_row += 1

    # Print each unique invoice and date
    for inv in invoices_list:
        c_inv = ws.cell(row=current_row, column=2, value=inv['invoice_no'])
        c_inv.alignment = Alignment(horizontal="center")

        c_dt = ws.cell(row=current_row, column=3, value=inv['date'])
        c_dt.alignment = Alignment(horizontal="center")
        current_row += 1

    # Set practical column widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 80
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 16
    ws.column_dimensions['G'].width = 16
    ws.column_dimensions['H'].width = 18

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer


# ---------------- Streamlit UI ----------------
st.title("📄 Invoice to Excel Generator")
st.write("Upload your invoice PDFs to extract, group by HSN & unit, isolate `+` categories, and generate the formatted Excel report.")

uploaded_files = st.file_uploader(
    "Upload Invoice PDFs",
    type=["pdf"],
    accept_multiple_files=True,
    help="Select all invoice PDFs to process."
)

if uploaded_files:
    st.info(f"Loaded {len(uploaded_files)} PDF file(s). Processing...")

    all_parsed_docs = []
    for file in uploaded_files:
        parsed = parse_pdf_stream(file)
        if parsed:
            parsed['filename'] = file.name
            all_parsed_docs.append(parsed)

    # Collect unique invoice numbers and dates
    invoices_list = []
    seen_invoices = set()
    for doc in all_parsed_docs:
        inv_no = doc.get('invoice_no', 'UNKNOWN')
        inv_date = doc.get('invoice_date', '')
        if inv_no not in seen_invoices:
            seen_invoices.add(inv_no)
            invoices_list.append({'invoice_no': inv_no, 'date': inv_date})
            
    # Consolidation across uploaded invoices
    consolidated = OrderedDict()
    for doc in all_parsed_docs:
        hsn = doc['hsn_code']
        base_desc = doc['base_description']
        conv_rate = doc.get('conv_rate', 1.0)
        # Extract KT from description (e.g. 18KT, 09KT) and extract the base prefix
        found_kts = re.findall(r'\b\d{1,2}KT\b', base_desc, flags=re.IGNORECASE)
        base_prefix = re.sub(r'\s*IN\s+.*$', '', base_desc, flags=re.IGNORECASE).strip()   
        for cat in doc['categories']:
            unit = cat['unit']
            header = cat['header']

            group_key = (hsn, unit, header) if cat['has_plus'] else (hsn, unit, "STANDARD")
            if group_key not in consolidated:
                consolidated[group_key] = {
                    'hsn': hsn,
                    'unit': unit,
                    'base_prefix': base_prefix,
                    'kts': set(found_kts),
                    'item_names': [],
                    'net_wt': 0.0,
                    'gross_wt': 0.0,
                    'qty': 0,
                    'amount': 0.0,
                    'amt_inr': 0.0
                }

            group = consolidated[group_key]
            group['kts'].update(found_kts)

            for it in cat['items']:
                if it['name'] not in group['item_names']:
                    group['item_names'].append(it['name'])
                group['net_wt'] += it.get('net_wt', 0.0)
                group['gross_wt'] += it['gross_wt']
                group['qty'] += int(it['qty'])
                group['amount'] += it['amount']
                group['amt_inr'] += round(it['amount'] * conv_rate, 2)

    # Display preview table
    preview_rows = []
    sr = 1
    for key, data in consolidated.items():
        kt_str = f" IN {', '.join(sorted(data['kts']))}" if data.get('kts') else ""
        base_description = f"{data.get('base_prefix', '')}{kt_str}"
        items_str = ", ".join(data['item_names'])
        desc_text = f"{base_description}-{items_str}-{data['qty']} {data['unit']}"

        preview_rows.append({
            "SR": sr,
            "HSN": data['hsn'],
            "DESCRIPTION": desc_text,
            "GrossWt": round(data['gross_wt'], 3),
            "NET WT": round(data.get('net_wt', 0.0), 3),       # <--- Added
            "QTY": data['qty'],
            "AMT": round(data['amount'], 2),
            "AMT INR": round(data.get('amt_inr', 0.0), 2)       # <--- Added
        })
        sr += 1

    df_preview = pd.DataFrame(preview_rows)
    st.subheader("Consolidated Preview")
    st.dataframe(df_preview, width='stretch')

    # Key Totals
    # Key Totals
    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("Total Gross Wt (Gms)", f"{df_preview['GrossWt'].sum():.3f}")
    col2.metric("Total Net Wt (Gms)", f"{df_preview['NET WT'].sum():.3f}")
    col3.metric("Total Quantity", f"{df_preview['QTY'].sum():,}")
    col4.metric("Total Amount (USD)", f"${df_preview['AMT'].sum():,.2f}")
    col5.metric("Total Amount (INR)", f"₹{df_preview['AMT INR'].sum():,.2f}")

    # Generate Excel download
    excel_file = generate_excel(consolidated, invoices_list)
    st.download_button(
        label="📥 Download Consolidated Excel File",
        data=excel_file,
        file_name="Consolidated_Invoice_Summary.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )